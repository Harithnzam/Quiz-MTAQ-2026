import sys
from docx import Document

path = sys.argv[1]
doc = Document(path)
for para in doc.paragraphs:
    if para.text.strip():
        print(para.text)
for t, table in enumerate(doc.tables):
    print(f"\n--- TABLE {t+1} ---")
    for row in table.rows:
        cells = [c.text.strip() for c in row.cells]
        print(" | ".join(cells))
