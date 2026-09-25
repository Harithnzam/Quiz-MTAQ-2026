import os
from docx import Document
import pdfplumber

os.makedirs("extracted", exist_ok=True)

# Extract all question banks
qb_dir = "Question Bank"
for fn in sorted(os.listdir(qb_dir)):
    if fn.endswith(".docx"):
        doc = Document(os.path.join(qb_dir, fn))
        out = []
        for para in doc.paragraphs:
            if para.text.strip():
                out.append(para.text)
        for t, table in enumerate(doc.tables):
            out.append(f"\n--- TABLE {t+1} ---")
            for row in table.rows:
                out.append(" | ".join(c.text.strip() for c in row.cells))
        outname = "extracted/" + fn.replace(".docx", ".txt")
        with open(outname, "w", encoding="utf-8") as f:
            f.write("\n".join(out))
        print(f"Wrote {outname} ({len(out)} blocks)")

print("DONE QB")
